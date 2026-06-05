"""Xây dựng lại KLTN v2 từ v1 — chỉnh tại chỗ, giữ nguyên hình ảnh/media.

Chiến lược:
  1. Copy v1 → v2 (giữ toàn bộ quan hệ media/ảnh).
  2. Xóa nội dung cũ của các mục cần thay (1.5, 1.6, Ch5) — không chèn thêm.
  3. Chèn mục mới (3.6, 4.6, 4.7) đúng vị trí.
  4. Cập nhật Tóm tắt, Mục lục, Tài liệu tham khảo.
  5. Định dạng Khoa + chuẩn hóa văn phong.

Usage: python scripts/build_thesis_v2.py
"""
from __future__ import annotations

import re
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

REPO = Path(__file__).resolve().parent.parent
V1 = REPO / "document" / "KLTN_TruongCongHuan_Final (1).docx"
OUT = REPO / "document" / "KLTN_TruongCongHuan_Final_v2.docx"

sys.path.insert(0, str(REPO / "thesis" / "content_v2"))
import sections as C  # noqa: E402

FONT = "Times New Roman"
BODY = Pt(13)
SECTION = Pt(14)
CHAPTER = Pt(16)

# Paragraph indices in v1 (verified)
IDX = {
    "sec_1_5": 207,
    "sec_1_6": 213,
    "ch2_start": 223,
    "ch3_start": 317,
    "ch4_start": 475,
    "ch5_start": 582,
    "refs_start": 615,
}


BODY_START = 175  # first CHƯƠNG 1 body paragraph in v1


def find_para(doc: Document, pattern: str, start: int = 0, body_only: bool = False) -> int | None:
    lo = BODY_START if body_only else start
    for i in range(lo, len(doc.paragraphs)):
        if pattern.lower() in doc.paragraphs[i].text.lower():
            return i
    return None


def find_section_heading(doc: Document, section_id: str) -> int | None:
    """Match exact section heading like '4.6.' at start (not '2.4.6.')."""
    sid = section_id.strip().rstrip(".")
    for i in range(BODY_START, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip()
        if re.match(rf"^{re.escape(sid)}\.", t) or t.startswith(section_id):
            return i
    return None


def _find_refs_heading(doc: Document) -> int | None:
    for i in range(BODY_START, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip().upper()
        if t == "TÀI LIỆU THAM KHẢO" or t.startswith("TÀI LIỆU THAM KHẢO") and len(t) < 30:
            return i
    return None


def find_chapter_heading(doc: Document, chapter_num: int) -> int | None:
    """Find exact 'CHƯƠNG N.' heading in body (not in-text mentions)."""
    prefix = f"CHƯƠNG {chapter_num}"
    for i in range(BODY_START, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip()
        if t.upper().startswith(prefix) and len(t) < 80:
            return i
    return None


def find_next_section(doc: Document, start: int, prefix: str) -> int:
    """Find next heading at same or higher level after start."""
    for i in range(start + 1, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip()
        if t.startswith(prefix):
            return i
        if t.startswith("CHƯƠNG"):
            return i
    return len(doc.paragraphs)


def delete_paragraphs(doc: Document, start: int, end: int) -> None:
    """Remove paragraphs [start, end) by deleting XML elements."""
    body = doc.element.body
    to_remove = []
    for i in range(start, min(end, len(doc.paragraphs))):
        el = doc.paragraphs[i]._element
        to_remove.append(el)
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def fmt_para(p, size=BODY, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.left_indent = None
    pf.first_line_indent = None
    text = p.text
    if not p.runs and text:
        p.add_run(text)
    for run in p.runs:
        run.font.name = FONT
        run.font.size = size
        if bold:
            run.bold = True


def insert_paragraphs_after(doc: Document, after_idx: int, blocks: list[str]) -> int:
    """Insert paragraphs after given index. Returns count inserted."""
    if after_idx < 0 or after_idx >= len(doc.paragraphs):
        ref = doc.paragraphs[-1]._element
    else:
        ref = doc.paragraphs[after_idx]._element

    inserted = 0
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        parts = block.split("\n\n")
        for part in parts:
            part = part.strip()
            if not part:
                continue
            first_line = part.split("\n")[0].strip()
            new_p = OxmlElement("w:p")
            ref.addnext(new_p)
            ref = new_p
            # Find the paragraph object
            from docx.text.paragraph import Paragraph
            para = Paragraph(new_p, doc)
            para.add_run(part.replace("\n", " ") if not first_line.startswith(("(", "-")) else part)

            if first_line.startswith("CHƯƠNG"):
                fmt_para(para, CHAPTER, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            elif re.match(r"^\d+\.\d+\.\d+\.", first_line):
                fmt_para(para, BODY, bold=True)
            elif re.match(r"^\d+\.\d+\.", first_line) and len(first_line) < 100:
                fmt_para(para, SECTION, bold=True)
            else:
                fmt_para(para, BODY)
            inserted += 1
    return inserted


def text_to_blocks(text: str) -> list[str]:
    return [b.strip() for b in text.strip().split("\n\n") if b.strip()]


def replace_section(doc: Document, heading_pattern: str, new_content: str, next_prefix: str) -> None:
    """Delete old section body and insert new content after heading."""
    idx = find_para(doc, heading_pattern, body_only=True)
    if idx is None:
        print(f"  [WARN] Không tìm thấy: {heading_pattern}")
        return
    end = find_next_section(doc, idx, next_prefix)
    # Keep heading, delete content until next section
    delete_paragraphs(doc, idx + 1, end)
    # Re-find idx after deletion
    idx = find_para(doc, heading_pattern, body_only=True)
    blocks = text_to_blocks(new_content)
    # First block is heading — skip if duplicate
    if blocks and blocks[0].startswith(heading_pattern.split(".")[0]):
        blocks = blocks[1:] if len(blocks) > 1 else []
    insert_paragraphs_after(doc, idx, blocks)
    print(f"  ✓ Thay thế {heading_pattern.strip()} (xóa {end - idx - 1} đoạn cũ)")


def insert_new_section(doc: Document, after_pattern: str, content: str) -> None:
    """Insert entirely new section after the paragraph matching pattern."""
    idx = find_para(doc, after_pattern, body_only=True)
    if idx is None:
        print(f"  [WARN] Không tìm thấy vị trí chèn sau: {after_pattern}")
        return
    # Find end of that section (last para before next CHƯƠNG or same-level heading)
    end = idx
    for i in range(idx + 1, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip()
        if t.startswith("CHƯƠNG"):
            break
        if re.match(r"^\d+\.\d+\.", t) and not t.startswith(after_pattern[:4]):
            # same level new section
            if not t.startswith(after_pattern.split()[0]):
                break
        end = i
    blocks = text_to_blocks(content)
    insert_paragraphs_after(doc, end, blocks)
    print(f"  ✓ Chèn mới sau {after_pattern.strip()}")


def update_abstract(doc: Document) -> None:
    idx = find_para(doc, "TÓM TẮT KHÓA LUẬN")
    if idx is None:
        idx = find_para(doc, "TÓM TẮT")
    if idx is None:
        return
    for j in range(idx + 1, min(idx + 8, len(doc.paragraphs))):
        p = doc.paragraphs[j]
        t = p.text.strip().upper()
        if any(x in t for x in ["MỤC LỤC", "DANH MỤC"]) or t.startswith("CHƯƠNG"):
            break
        if p.text.strip():
            for run in p.runs:
                run.text = ""
            p.add_run(C.ABSTRACT_V2)
            fmt_para(p, BODY)
            print("  ✓ Cập nhật Tóm tắt")
            return


def update_toc(doc: Document) -> None:
    """Add new TOC lines in front-matter TOC (before body)."""
    for extra in C.TOC_EXTRA_LINES:
        # Check TOC area only (before BODY_START)
        found_in_toc = any(
            extra[:25] in doc.paragraphs[i].text
            for i in range(min(BODY_START, len(doc.paragraphs)))
        )
        if found_in_toc:
            continue
        idx = None
        for i in range(min(BODY_START, len(doc.paragraphs))):
            if "4.5. Đánh giá hệ thống" in doc.paragraphs[i].text:
                idx = i
                break
        if idx is None:
            continue
        p = doc.paragraphs[idx]
        ref = p._element
        new_p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        t_el = OxmlElement("w:t")
        short = extra[:55]
        t_el.text = short + "…" * max(1, 60 - len(short))
        r.append(t_el)
        new_p.append(r)
        ref.addnext(new_p)
        print(f"  ✓ Thêm vào Mục lục: {extra[:40]}…")


def append_refs(doc: Document) -> None:
    idx = _find_refs_heading(doc)
    if idx is None:
        return
    existing = doc.paragraphs[idx:].__len__()
    for ref in C.NEW_REFS:
        if find_para(doc, ref[:15]) is not None:
            continue
        p = doc.add_paragraph(ref)
        fmt_para(p, BODY)
    print("  ✓ Bổ sung tài liệu tham khảo [20]–[24]")


def apply_margins(doc: Document) -> None:
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(1.5)


def apply_global_format(doc: Document) -> None:
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5
        if pf.left_indent and pf.left_indent > Cm(1.5):
            pf.left_indent = None
        if t.startswith("CHƯƠNG") and len(t) < 80:
            fmt_para(p, CHAPTER, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif re.match(r"^\d+\.\d+\.", t) and len(t) < 100 and not t.endswith("."):
            fmt_para(p, SECTION, bold=True)
        else:
            for run in p.runs:
                if run.font.name is None or run.font.name == "Calibri":
                    run.font.name = FONT
                if not run.font.size:
                    run.font.size = BODY

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = FONT
                        run.font.size = Pt(11)


def _run_has_image(run) -> bool:
    xml = run._element.xml
    return "blip" in xml or "imagedata" in xml or "drawing" in xml or "pict" in xml


def terminology_pass(doc: Document) -> int:
    repl = {
        r"\bTLU AI Assistant\b": "TLU Assistant System",
        r"\btôi đã\b": "tác giả đã",
        r"\btôi thực hiện\b": "tác giả thực hiện",
        r"\btôi thẳng thắn\b": "khóa luận thẳng thắn",
        r"\btôi đề xuất\b": "khóa luận đề xuất",
        r"\bTôi xin\b": "Tác giả xin",
        r"\bLời đầu tiên, tôi\b": "Lời đầu tiên, tác giả",
        r"\bCuối cùng, tôi\b": "Cuối cùng, tác giả",
    }
    n = 0
    for p in doc.paragraphs:
        for run in p.runs:
            if _run_has_image(run) or not run.text:
                continue
            for pat, rep in repl.items():
                run.text, c = re.subn(pat, rep, run.text, flags=re.IGNORECASE)
                n += c
    return n


def remove_duplicate_sections(doc: Document) -> None:
    """Remove duplicate 3.6/4.6/4.7 blocks left by old patch script (body only)."""
    seen: dict[str, int] = {}
    to_remove = []
    for i in range(BODY_START, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip()
        for key in ["3.6. Ánh xạ", "4.6. Đánh giá định lượng", "4.7. Quan sát"]:
            if t.startswith(key):
                if key in seen:
                    end = find_next_section(doc, i, "3." if "3.6" in key else "4.")
                    for j in range(i, end):
                        to_remove.append(doc.paragraphs[j]._element)
                else:
                    seen[key] = i
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
    if to_remove:
        print(f"  ✓ Xóa {len(to_remove)} đoạn trùng lặp (body)")


def replace_chapter5(doc: Document) -> None:
    idx = find_chapter_heading(doc, 5)
    if idx is None:
        print("  [WARN] Không tìm thấy tiêu đề CHƯƠNG 5")
        return
    refs_idx = _find_refs_heading(doc)
    if refs_idx is None:
        refs_idx = len(doc.paragraphs)
    delete_paragraphs(doc, idx + 1, refs_idx)
    idx = find_chapter_heading(doc, 5)
    all_blocks = (
        text_to_blocks(C.CH5_1)
        + text_to_blocks(C.CH5_2)
        + text_to_blocks(C.CH5_3)
    )
    insert_paragraphs_after(doc, idx, all_blocks)
    print("  ✓ Viết lại toàn bộ Chương 5")


def build() -> None:
    if not V1.exists():
        print(f"[ERROR] Không tìm thấy {V1}")
        sys.exit(1)

    print(f"[1] Copy {V1.name} → {OUT.name}")
    shutil.copy2(V1, OUT)

    print("[2] Mở và dọn dẹp bản patch cũ (nếu có)")
    doc = Document(str(OUT))
    remove_duplicate_sections(doc)

    print("[3] Thay thế / chèn nội dung mới theo khung AI Engineering")
    replace_section(doc, "1.5. Đóng góp của khóa luận", C.SEC_1_5, "1.6")
    replace_section(doc, "1.6. Cấu trúc khóa luận", C.SEC_1_6, "CHƯƠNG 2")

    # 3.6 — insert after 3.5 if not exists in body
    if find_section_heading(doc, "3.6") is None:
        insert_new_section(doc, "3.5.2. Giao diện Staff Console", C.SEC_3_6)

    # 4.6, 4.7 — insert after 4.5 if not exists in body
    if find_section_heading(doc, "4.6") is None:
        insert_new_section(doc, "4.5.2. Điểm mạnh", C.SEC_4_6)
    if find_section_heading(doc, "4.7") is None:
        idx46 = find_section_heading(doc, "4.6")
        if idx46 is not None:
            insert_new_section(doc, "4.6. Đánh giá định lượng", C.SEC_4_7)
        else:
            insert_new_section(doc, "4.5.2. Điểm mạnh", C.SEC_4_7)

    replace_chapter5(doc)
    update_abstract(doc)
    update_toc(doc)
    append_refs(doc)

    print("[4] Định dạng Khoa + chuẩn hóa văn phong")
    apply_margins(doc)
    apply_global_format(doc)
    n = terminology_pass(doc)
    print(f"  ✓ Chuẩn hóa thuật ngữ: {n} chỗ")

    print(f"[5] Lưu {OUT}")
    doc.save(str(OUT))

    # Verify
    v = Document(str(OUT))
    dups = {}
    for i, p in enumerate(v.paragraphs):
        if i < BODY_START:
            continue
        for k in ["1.5. Đóng góp", "3.6. Ánh xạ", "4.6. Đánh giá", "4.7. Quan sát"]:
            if p.text.strip().startswith(k):
                dups[k] = dups.get(k, 0) + 1
    print(f"\nHoàn tất — {len(v.paragraphs)} đoạn, {len(v.tables)} bảng, {len(v.inline_shapes)} hình")
    for k, c in dups.items():
        status = "OK" if c == 1 else f"CẢNH BÁO: {c} lần"
        print(f"  {k}: {status}")
    print("\nMở Word → Ctrl+A → F9 để cập nhật Mục lục.")
    print("Chạy run_eval.py --ablation để thay số thật vào Bảng 4.5–4.7.")


if __name__ == "__main__":
    build()
