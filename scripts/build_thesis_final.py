#!/usr/bin/env python3
"""Tạo KLTN hoàn chỉnh từ đầu — không cần template.

Chuẩn định dạng: Khoa Trí tuệ Nhân tạo / CNTT — Đại học Thăng Long
  - Times New Roman, Unicode
  - 13pt body | 14pt bold section | 16pt bold chapter
  - Lề: trên/dưới 2cm, trái 3cm, phải 1.5cm
  - Giãn dòng 1.5; số trang Ả Rập căn giữa footer

Usage:
    python3 scripts/build_thesis_final.py

Output:
    document/KLTN_TruongCongHuan_Final.docx
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ── Paths ────────────────────────────────────────────────────────────────────
REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "thesis" / "content_v3"))
import chapters as C  # noqa: E402

OUT = REPO / "document" / "KLTN_TruongCongHuan_Final.docx"

# ── Typography constants ──────────────────────────────────────────────────────
FONT = "Times New Roman"
PT13 = Pt(13)
PT14 = Pt(14)
PT16 = Pt(16)


# ── Helpers ───────────────────────────────────────────────────────────────────

def _setup_margins(doc: Document) -> None:
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(1.5)


def _fmt(p,
         size: Pt = PT13,
         bold: bool = False,
         italic: bool = False,
         align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.JUSTIFY,
         space_before: Pt = Pt(0),
         space_after: Pt = Pt(6)) -> None:
    """Apply TLU paragraph formatting."""
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = space_before
    pf.space_after = space_after
    for run in p.runs:
        run.font.name = FONT
        run.font.size = size
        run.font.bold = bold
        run.font.italic = italic


def _para(doc: Document, text: str = "",
          size: Pt = PT13,
          bold: bool = False,
          italic: bool = False,
          align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.JUSTIFY,
          space_before: Pt = Pt(0),
          space_after: Pt = Pt(6)) -> None:
    p = doc.add_paragraph()
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = size
        run.font.bold = bold
        run.font.italic = italic
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = space_before
    pf.space_after = space_after


def _chapter(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.font.name = FONT
    run.font.size = PT16
    run.font.bold = True
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)


def _section(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = PT14
    run.font.bold = True
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)


def _subsection(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = PT13
    run.font.bold = True
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)


def _body(doc: Document, text: str) -> None:
    _para(doc, text)


def _bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = PT13
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)


def _add_table(doc: Document, caption: str, headers: list[str],
               rows: list[list[str]]) -> None:
    _para(doc, caption, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
          space_before=Pt(6), space_after=Pt(3))
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        tbl.style = "Table Grid"
    except KeyError:
        pass
    # Header row
    for ci, h in enumerate(headers):
        cell = tbl.rows[0].cells[ci]
        cell.text = h
        for cp in cell.paragraphs:
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in cp.runs:
                r.font.name = FONT
                r.font.size = Pt(11)
                r.bold = True
    # Data rows
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for cp in cell.paragraphs:
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in cp.runs:
                    r.font.name = FONT
                    r.font.size = Pt(11)
    _para(doc, "")


def _page_break(doc: Document) -> None:
    doc.add_page_break()


def _add_footer_page_number(doc: Document) -> None:
    """Số trang Ả Rập, căn giữa, trong lề dưới."""
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        if footer.paragraphs:
            para = footer.paragraphs[0]
        else:
            para = footer.add_paragraph()
        para.clear()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        run.font.name = FONT
        run.font.size = PT13

        fc_b = OxmlElement("w:fldChar")
        fc_b.set(qn("w:fldCharType"), "begin")
        run._r.append(fc_b)

        ir = para.add_run()
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        ir._r.append(instr)

        er = para.add_run()
        fc_e = OxmlElement("w:fldChar")
        fc_e.set(qn("w:fldCharType"), "end")
        er._r.append(fc_e)


# ══════════════════════════════════════════════════════════════════════════════
#  COVER PAGES
# ══════════════════════════════════════════════════════════════════════════════

def _add_cover(doc: Document) -> None:
    """Trang bìa chính."""
    _para(doc, "BỘ GIÁO DỤC VÀ ĐÀO TẠO", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(0))
    _para(doc, "TRƯỜNG ĐẠI HỌC THĂNG LONG", size=PT14, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(24))
    _para(doc, "─────────────────────────────",
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(24))

    _para(doc, "KHÓA LUẬN TỐT NGHIỆP", size=PT16, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
    _para(doc, "Ngành: Khoa học Máy tính", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(30))

    _para(doc, "Đề tài:", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))
    _para(doc, "PHÁT TRIỂN HỆ THỐNG TRỢ LÝ AI HỌC VỤ",
          size=PT16, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))
    _para(doc, "TÍCH HỢP QUẢN LÝ SINH VIÊN",
          size=PT16, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(36))

    _para(doc, "Sinh viên thực hiện:", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))
    _para(doc, "TRƯƠNG CÔNG HUẤN", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))

    _para(doc, "Giảng viên hướng dẫn:", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))
    _para(doc, "ThS. NGÔ MẠNH CƯỜNG", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(36))

    _para(doc, "Hà Nội, 2026", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(0))

    _page_break(doc)


def _add_phubìa(doc: Document) -> None:
    """Trang phụ bìa."""
    _para(doc, "BỘ GIÁO DỤC VÀ ĐÀO TẠO", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(0))
    _para(doc, "TRƯỜNG ĐẠI HỌC THĂNG LONG", size=PT14, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
    _para(doc, "KHOA CÔNG NGHỆ THÔNG TIN", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(24))
    _para(doc, "─────────────────────────────",
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(24))

    _para(doc, "KHÓA LUẬN TỐT NGHIỆP", size=PT16, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
    _para(doc, "Ngành: Khoa học Máy tính", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(30))

    _para(doc, "Đề tài:", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))
    _para(doc, "PHÁT TRIỂN HỆ THỐNG TRỢ LÝ AI HỌC VỤ",
          size=PT16, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))
    _para(doc, "TÍCH HỢP QUẢN LÝ SINH VIÊN",
          size=PT16, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(30))

    _para(doc, "Sinh viên thực hiện: TRƯƠNG CÔNG HUẤN", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))
    _para(doc, "Mã sinh viên: A33941", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))
    _para(doc, "Chuyên ngành: Khoa học Máy tính", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
    _para(doc, "Giảng viên hướng dẫn: ThS. NGÔ MẠNH CƯỜNG", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(36))

    _para(doc, "Hà Nội, 2026", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(0))

    _page_break(doc)


# ══════════════════════════════════════════════════════════════════════════════
#  FRONT MATTER
# ══════════════════════════════════════════════════════════════════════════════

def _add_loi_cam_on(doc: Document) -> None:
    _chapter(doc, "LỜI CẢM ƠN")
    _body(doc,
          "Trước hết, em xin gửi lời cảm ơn chân thành tới ThS. Ngô Mạnh Cường — người đã "
          "tận tình hướng dẫn, góp ý và tạo điều kiện thuận lợi để em hoàn thành khóa luận "
          "tốt nghiệp với đề tài Phát triển hệ thống trợ lý AI học vụ tích hợp quản lý "
          "sinh viên tại Trường Đại học Thăng Long.")
    _body(doc,
          "Em xin cảm ơn các thầy cô Khoa Công nghệ Thông tin — Trường Đại học Thăng Long "
          "đã truyền đạt kiến thức nền tảng về công nghệ thông tin, trí tuệ nhân tạo, phân "
          "tích thiết kế hệ thống và phương pháp nghiên cứu khoa học trong suốt quá trình "
          "học tập.")
    _body(doc,
          "Em cũng xin gửi lời cảm ơn tới Ban Giám hiệu, các phòng ban và cán bộ nhà "
          "trường đã cung cấp thông tin nghiệp vụ, tài liệu quy chế tham khảo khi xây dựng "
          "kịch bản hệ thống. Đồng thời, em xin cảm ơn gia đình và bạn bè đã động viên, "
          "hỗ trợ tinh thần trong suốt thời gian thực hiện đề tài.")
    _body(doc,
          "Do thời gian và kinh nghiệm còn hạn chế, khóa luận không tránh khỏi thiếu sót. "
          "Em rất mong nhận được sự góp ý của quý thầy cô và hội đồng chấm.")
    _body(doc, "Em xin chân thành cảm ơn!")
    # Blank lines before signature
    for _ in range(5):
        _para(doc, "")
    # Thông tin sinh viên: 2 dòng căn phải theo quy định
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run("SINH VIÊN")
    r1.font.name = FONT
    r1.font.size = PT13
    r1.font.bold = True
    pf = p1.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(65)   # 65pt gap theo quy định

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run("Trương Công Huấn")
    r2.font.name = FONT
    r2.font.size = PT13
    r2.font.bold = True
    pf2 = p2.paragraph_format
    pf2.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf2.line_spacing = 1.5
    pf2.space_before = Pt(0)
    pf2.space_after = Pt(0)

    _page_break(doc)


def _add_loi_cam_doan(doc: Document) -> None:
    _chapter(doc, "LỜI CAM ĐOAN")
    _body(doc,
          "Em xin cam đoan khóa luận tốt nghiệp Phát triển hệ thống trợ lý AI học vụ tích "
          "hợp quản lý sinh viên tại Trường Đại học Thăng Long là công trình nghiên cứu do "
          "em thực hiện dưới sự hướng dẫn của ThS. Ngô Mạnh Cường. Các số liệu, kết quả "
          "thử nghiệm và nhận xét trong khóa luận là trung thực; các nội dung trích dẫn từ "
          "tài liệu của tác giả khác đều được ghi rõ nguồn theo quy định của nhà trường và "
          "tuân thủ Luật Sở hữu trí tuệ.")
    _body(doc,
          "Em cam đoan không sao chép, không sử dụng trái phép sản phẩm của người khác. "
          "Nếu có vi phạm, em xin chịu trách nhiệm trước nhà trường và pháp luật.")
    _body(doc, "Em xin trân trọng cảm ơn!")
    for _ in range(5):
        _para(doc, "")

    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run("SINH VIÊN")
    r1.font.name = FONT
    r1.font.size = PT13
    r1.font.bold = True
    pf = p1.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(65)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run("Trương Công Huấn")
    r2.font.name = FONT
    r2.font.size = PT13
    r2.font.bold = True
    pf2 = p2.paragraph_format
    pf2.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf2.line_spacing = 1.5
    pf2.space_before = Pt(0)
    pf2.space_after = Pt(0)

    _page_break(doc)


def _add_abstract(doc: Document) -> None:
    _chapter(doc, "TÓM TẮT KHÓA LUẬN")
    for para in C.ABSTRACT.strip().split("\n\n"):
        para = para.strip()
        if para:
            _body(doc, para)
    _page_break(doc)


def _add_muc_luc(doc: Document) -> None:
    _chapter(doc, "MỤC LỤC")
    toc_entries = [
        ("LỜI CẢM ƠN", False),
        ("LỜI CAM ĐOAN", False),
        ("TÓM TẮT KHÓA LUẬN", False),
        ("MỤC LỤC", False),
        ("DANH MỤC CÁC CHỮ VIẾT TẮT", False),
        ("DANH MỤC BẢNG BIỂU", False),
        ("MỞ ĐẦU", False),
        ("CHƯƠNG 1. GIỚI THIỆU", True),
        ("1.1. Phát biểu bài toán", False),
        ("1.2. Mục tiêu nghiên cứu", False),
        ("1.3. Phạm vi nghiên cứu", False),
        ("    1.3.1. Phạm vi chức năng", False),
        ("    1.3.2. Phạm vi kỹ thuật", False),
        ("1.4. Phương pháp tiếp cận", False),
        ("1.5. Đóng góp của khóa luận", False),
        ("1.6. Cấu trúc khóa luận", False),
        ("CHƯƠNG 2. CƠ SỞ LÝ THUYẾT VÀ CÁC PHƯƠNG PHÁP LIÊN QUAN", True),
        ("2.1. Cơ sở lý thuyết", False),
        ("    2.1.1. Trí tuệ nhân tạo và học máy", False),
        ("    2.1.2. Mô hình ngôn ngữ lớn", False),
        ("2.2. Các phương pháp tiếp cận hiện có", False),
        ("2.3. Các mô hình AI liên quan", False),
        ("    2.3.1. Embedding và vector search", False),
        ("    2.3.2. Retrieval-Augmented Generation", False),
        ("2.4. Hướng tiếp cận của khóa luận", False),
        ("CHƯƠNG 3. PHƯƠNG PHÁP ĐỀ XUẤT", True),
        ("3.1. Kiến trúc hệ thống", False),
        ("3.2. Dataset sử dụng", False),
        ("3.3. Tiền xử lý dữ liệu", False),
        ("    3.3.1. Trích xuất văn bản", False),
        ("    3.3.2. Semantic Markdown Chunking", False),
        ("3.4. Mô hình AI và thuật toán", False),
        ("    3.4.1. Embedding và vector store", False),
        ("    3.4.2. Query Expansion và retrieval", False),
        ("    3.4.3. Generation", False),
        ("3.5. Chiến lược indexing và cấu hình pipeline", False),
        ("CHƯƠNG 4. THỰC NGHIỆM VÀ ĐÁNH GIÁ", True),
        ("4.1. Môi trường thực nghiệm", False),
        ("4.2. Cấu hình pipeline RAG", False),
        ("4.3. Chỉ số đánh giá", False),
        ("4.4. Kết quả thực nghiệm", False),
        ("4.5. So sánh với các phương pháp khác", False),
        ("4.6. Phân tích kết quả", False),
        ("CHƯƠNG 5. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", True),
        ("5.1. Kết luận", False),
        ("5.2. Hạn chế của nghiên cứu", False),
        ("5.3. Hướng phát triển trong tương lai", False),
        ("TÀI LIỆU THAM KHẢO", False),
        ("PHỤ LỤC", False),
    ]
    for text, is_chap in toc_entries:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = PT13
        run.font.bold = is_chap
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
    _page_break(doc)


def _add_abbreviations(doc: Document) -> None:
    _chapter(doc, "DANH MỤC CÁC CHỮ VIẾT TẮT")
    tbl = doc.add_table(rows=1 + len(C.ABBREVIATIONS), cols=2)
    try:
        tbl.style = "Table Grid"
    except KeyError:
        pass
    for ci, hdr in enumerate(["Chữ viết tắt", "Thuật ngữ"]):
        cell = tbl.rows[0].cells[ci]
        cell.text = hdr
        for cp in cell.paragraphs:
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in cp.runs:
                r.font.name = FONT
                r.font.size = Pt(11)
                r.bold = True
    for ri, (abbr, full) in enumerate(C.ABBREVIATIONS):
        row = tbl.rows[ri + 1]
        row.cells[0].text = abbr
        row.cells[1].text = full
        for ci in range(2):
            for cp in row.cells[ci].paragraphs:
                for r in cp.runs:
                    r.font.name = FONT
                    r.font.size = Pt(11)
    _para(doc, "")
    _page_break(doc)


def _add_danh_muc_bang(doc: Document) -> None:
    _chapter(doc, "DANH MỤC BẢNG BIỂU")
    entries = [
        "Bảng 3.1. Phân loại cấu hình pipeline RAG trong ablation",
        "Bảng 4.3. Phân loại bộ câu hỏi đánh giá RAG (50 câu)",
        "Bảng 4.4. Kết quả ablation pipeline RAG trên corpus Sổ tay K38",
    ]
    for e in entries:
        _body(doc, e)
    _page_break(doc)


# ══════════════════════════════════════════════════════════════════════════════
#  MỞ ĐẦU
# ══════════════════════════════════════════════════════════════════════════════

def _add_mo_dau(doc: Document) -> None:
    _chapter(doc, "MỞ ĐẦU")

    def _sec(t: str) -> None:
        _section(doc, t)

    _sec("Bối cảnh nghiên cứu")
    _body(doc,
          "Giáo dục đại học Việt Nam đang đẩy mạnh chuyển đổi số: đăng ký học phần trực "
          "tuyến, E-Learning, thông báo điện tử và quản lý hồ sơ sinh viên trên các hệ "
          "thống rời rạc. Sinh viên thường phải tra cứu quy chế, quy trình nộp đơn và liên "
          "hệ từng phòng ban (cố vấn học tập, công tác sinh viên, đào tạo, tài chính, tâm "
          "lý) qua nhiều kênh — email, điện thoại, fanpage — dẫn đến chậm trễ và thiếu "
          "nhất quán thông tin. Song song đó, các mô hình ngôn ngữ lớn (LLM) và kỹ thuật "
          "Retrieval-Augmented Generation (RAG) cho phép hệ thống trả lời câu hỏi dựa trên "
          "tài liệu nội bộ của trường thay vì tạo nội dung không có căn cứ, phù hợp với "
          "bài toán hỏi đáp quy chế và hướng dẫn học vụ [1], [2].")
    _body(doc,
          "Trường Đại học Thăng Long (TLU) là môi trường triển khai thực nghiệm điển hình: "
          "sinh viên cần một điểm chạm thống nhất trên điện thoại; các phòng ban cần "
          "dashboard quản lý chat, thông báo, đơn từ và nguồn tri thức phục vụ trợ lý AI. "
          "Xuất phát từ yêu cầu đó, khóa luận hướng tới xây dựng TLU AI Assistant — sản "
          "phẩm phần mềm tích hợp trợ lý AI và kênh tương tác sinh viên–cán bộ.")

    _sec("Lý do chọn đề tài")
    _bullet(doc, "Tính thời sự: Ứng dụng AI generative và RAG trong giáo dục là xu hướng "
            "được quan tâm trong ngành CNTT và hệ thống thông tin [3].")
    _bullet(doc, "Tính thực tiễn: Đề tài bám sát nghiệp vụ thật tại TLU (năm phòng ban, "
            "thông báo, đơn từ, tài liệu quy chế).")
    _bullet(doc, "Tính kỹ thuật: Monorepo kết hợp mobile, web, BFF, microservice RAG và "
            "Firebase — phù hợp đánh giá năng lực tốt nghiệp ngành CNTT.")
    _bullet(doc, "Khả năng triển khai: Mã nguồn đã được hiện thực, có Docker Compose và "
            "checklist E2E, thuận lợi cho chương thực nghiệm.")

    _sec("Mục tiêu nghiên cứu")
    _bullet(doc, "Khảo sát và làm rõ yêu cầu chức năng, phi chức năng của hệ thống trợ lý "
            "AI và tương tác phòng ban đối với sinh viên TLU.")
    _bullet(doc, "Thiết kế kiến trúc tổng thể, mô hình dữ liệu và các luồng nghiệp vụ "
            "chính (RAG, chat, thông báo, đơn từ).")
    _bullet(doc, "Xây dựng prototype: ứng dụng Flutter, Staff Console, BFF Express, RAG "
            "API FastAPI, tích hợp Firestore và triển khai Docker.")
    _bullet(doc, "Kiểm thử, đánh giá mức độ đáp ứng yêu cầu và đề xuất hướng phát triển.")

    _sec("Đối tượng và phạm vi nghiên cứu")
    _body(doc,
          "Đối tượng: Quy trình tra cứu tri thức học vụ/quy chế bằng AI và quy trình tương "
          "tác sinh viên–phòng ban trên nền tảng số.")
    _body(doc,
          "Phạm vi chức năng: RAG (local/server), STT/TTS tiếng Việt, chat năm phòng ban, "
          "thông báo FCM, nộp đơn (Storage), góp ý, WebView cổng học vụ; phía cán bộ: quản "
          "lý SV, chat, thông báo, đơn, góp ý, tài liệu vector, audit, RBAC.")
    _body(doc,
          "Ngoài phạm vi: Thay thế hoàn toàn các hệ thống học vụ lõi (SIS, LMS); xử lý "
          "thanh toán học phí; nghiên cứu huấn luyện LLM từ đầu.")

    _sec("Phương pháp nghiên cứu")
    _bullet(doc, "Khảo sát, phân tích: Tổng hợp yêu cầu từ quy trình thực tế và tài liệu trường.")
    _bullet(doc, "Phân tích – thiết kế hướng đối tượng: Use case, sequence diagram, mô hình dữ liệu Firestore.")
    _bullet(doc, "Phát triển lặp: Xây dựng từng module (RAG → Firebase → chat → staff console), kiểm thử tích hợp theo checklist E2E.")
    _bullet(doc, "Thực nghiệm, đánh giá: Chạy kịch bản demo và đối chiếu tiêu chí chấp nhận.")

    _sec("Ý nghĩa khoa học và thực tiễn")
    _body(doc,
          "Về khoa học: Áp dụng pipeline RAG (chunking, embedding, retrieval, generation) "
          "trong miền dữ liệu giáo dục; so sánh triển khai on-device và server-side; "
          "thực nghiệm ablation định lượng với 12 cấu hình và bộ eval 50 câu hỏi.")
    _body(doc,
          "Về thực tiễn: Cung cấp giải pháp tham chiếu cho TLU và các trường tương tự khi "
          "triển khai trợ lý AI nội bộ kèm kênh hành chính số.")

    _sec("Cấu trúc khóa luận")
    _bullet(doc, "Chương 1 trình bày sự cần thiết, mục tiêu, phạm vi, phương pháp và đóng góp đề tài.")
    _bullet(doc, "Chương 2 trình bày cơ sở lý thuyết, công nghệ và phương pháp tiếp cận liên quan.")
    _bullet(doc, "Chương 3 mô tả phương pháp đề xuất: kiến trúc, dataset, tiền xử lý, mô hình AI.")
    _bullet(doc, "Chương 4 trình bày thực nghiệm và đánh giá kết quả.")
    _bullet(doc, "Chương 5 tổng kết kết quả, hạn chế và hướng phát triển.")
    _body(doc, "Khóa luận kết thúc bằng tài liệu tham khảo (chuẩn IEEE) và phụ lục.")
    _page_break(doc)


# ══════════════════════════════════════════════════════════════════════════════
#  BODY CHAPTERS (from chapters.py BODY_BLOCKS)
# ══════════════════════════════════════════════════════════════════════════════

def _add_chapters(doc: Document) -> None:
    table_map = {
        "EVAL_CATEGORY_TABLE": C.EVAL_CATEGORY_TABLE,
        "ABLACTION_TABLE": C.ABLACTION_TABLE,
    }

    kind_map = {
        "chapter": _chapter,
        "section": _section,
        "subsection": _subsection,
        "body": _body,
    }

    for kind, text in C.BODY_BLOCKS:
        if kind == "table_ref":
            spec = table_map[text]
            _add_table(doc, spec["caption"], spec["headers"], spec["rows"])
        elif kind == "bullet":
            _bullet(doc, text)
        elif kind in kind_map:
            kind_map[kind](doc, text)

    _page_break(doc)


# ══════════════════════════════════════════════════════════════════════════════
#  REFERENCES
# ══════════════════════════════════════════════════════════════════════════════

def _add_references(doc: Document) -> None:
    _chapter(doc, "TÀI LIỆU THAM KHẢO")
    for ref in C.REFERENCES:
        p = doc.add_paragraph()
        run = p.add_run(ref)
        run.font.name = FONT
        run.font.size = PT13
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        # Thêm indent treo cho trích dẫn IEEE [n]
        pf.left_indent = Pt(24)
        pf.first_line_indent = Pt(-24)
    _page_break(doc)


# ══════════════════════════════════════════════════════════════════════════════
#  PHỤ LỤC
# ══════════════════════════════════════════════════════════════════════════════

def _add_phu_luc(doc: Document) -> None:
    _chapter(doc, "PHỤ LỤC")

    # Phụ lục 1 — Danh mục phòng ban
    _section(doc, "Phụ lục 1. Danh mục phòng ban và mã định danh")
    dept_headers = ["STT", "Mã (departmentId)", "Tên phòng ban"]
    dept_rows = [
        ["1", "co_van", "Cố vấn học tập"],
        ["2", "ct_sv", "Phòng Công tác Chính trị Sinh viên"],
        ["3", "dao_tao", "Phòng Đào tạo"],
        ["4", "tai_chinh", "Phòng Tài chính – Kế toán"],
        ["5", "tam_ly", "Phòng Tham vấn tâm lý"],
    ]
    _add_table(doc, "", dept_headers, dept_rows)

    # Phụ lục 2 — Biến môi trường
    _section(doc, "Phụ lục 2. Biến môi trường mẫu (.env.example)")
    env_headers = ["Biến", "Mô tả"]
    env_rows = [
        ["GEMINI_API_KEY", "Khóa Google Gemini (embedding + generate)"],
        ["STUDENT_API_KEY", "Khóa xác thực sinh viên gọi RAG API"],
        ["ADMIN_API_KEY", "Khóa quản trị ingest tài liệu RAG"],
        ["STAFF_API_KEY", "Khóa bypass vận hành BFF (không public)"],
        ["JWT_SECRET", "Secret ký JWT cán bộ"],
        ["FIREBASE_SERVICE_ACCOUNT_JSON", "JSON service account (hoặc mount file)"],
        ["CHROMA_DATA_DIR", "Thư mục lưu ChromaDB persistent volume"],
        ["SOURCES_DIR", "Thư mục file tài liệu upload"],
        ["STORAGE_BACKEND", "file (Docker) hoặc firebase (production)"],
    ]
    _add_table(doc, "", env_headers, env_rows)

    _body(doc,
          "Flutter (--dart-define): GEMINI_KEY (Gemini RAG on-device), "
          "API_BASE_URL (URL RAG API, ví dụ http://10.0.2.2:8080), "
          "STUDENT_API_KEY (header X-Student-Key).")

    # Phụ lục 3 — Endpoint RAG API
    _section(doc, "Phụ lục 3. Endpoint RAG API (tóm tắt)")
    api_headers = ["Phương thức", "Đường dẫn", "Xác thực"]
    api_rows = [
        ["GET", "/health", "Public"],
        ["POST", "/admin/documents", "X-Admin-Key"],
        ["GET", "/admin/documents", "X-Admin-Key"],
        ["DELETE", "/admin/documents/{id}", "X-Admin-Key"],
        ["GET", "/admin/ingest-jobs/{id}", "X-Admin-Key"],
        ["POST", "/v1/chat/stream", "X-Student-Key (SSE)"],
    ]
    _add_table(doc, "", api_headers, api_rows)

    # Phụ lục 4 — Cấu trúc mã nguồn Flutter
    _section(doc, "Phụ lục 4. Cấu trúc thư mục mã nguồn Flutter")
    _body(doc,
          "lib/ — mã nguồn ứng dụng Flutter: "
          "main.dart (entry point); "
          "config/ (api_config, firebase, department_catalog, demo); "
          "models/ (dữ liệu); "
          "providers/ (ChatProvider, trạng thái UI); "
          "screens/ (màn hình chính); "
          "services/ (RAG local/remote, Firestore, TTS/STT); "
          "widgets/ (component tái sử dụng); "
          "theme/ (Dark theme, typography).")

    # Phụ lục 5 — Route Staff Console
    _section(doc, "Phụ lục 5. Route Staff Console")
    route_headers = ["Route", "Trang"]
    route_rows = [
        ["/login", "Đăng nhập"],
        ["/", "Tổng quan (Dashboard)"],
        ["/students", "Hồ sơ sinh viên"],
        ["/chat", "Chat phòng ban"],
        ["/notifications", "Thông báo"],
        ["/submissions", "Đơn từ"],
        ["/feedback", "Góp ý"],
        ["/vectors", "Tài liệu RAG"],
        ["/audit", "Nhật ký kiểm toán"],
        ["/settings", "Cài đặt"],
    ]
    _add_table(doc, "", route_headers, route_rows)

    # Phụ lục 6 — Hướng dẫn triển khai server Docker
    _section(doc, "Phụ lục 6. Hướng dẫn triển khai server Docker")
    _body(doc,
          "Yêu cầu: Docker Engine ≥ 24, Docker Compose Plugin, GEMINI_API_KEY hợp lệ.")
    _body(doc, "Các bước triển khai:")
    steps = [
        "git clone <repo_url> && cd Graduation_Thesis",
        "cp .env.example .env  # Điền GEMINI_API_KEY, STUDENT_API_KEY, ADMIN_API_KEY, STAFF_API_KEY, JWT_SECRET",
        "# Đặt file Firebase service account vào secrets/firebase-sa.json",
        "docker compose up --build -d",
        "# Kiểm tra: curl http://127.0.0.1:8080/health",
        "# Staff Console: http://127.0.0.1:3000 — đăng nhập super_admin / 123456",
        "# Ingest corpus: POST /admin/documents với X-Admin-Key qua Staff Console /vectors",
    ]
    for step in steps:
        _bullet(doc, step)

    _body(doc,
          "Cổng dịch vụ: RAG API tại :8080, Staff Console (UI + API) tại :3000. "
          "Flutter app kết nối qua --dart-define=API_BASE_URL=http://<IP_SERVER>:8080. "
          "Xem thêm: docs/DEPLOY_CLOUD_RUN.md để triển khai lên Google Cloud Run.")


# ══════════════════════════════════════════════════════════════════════════════
#  MAIN
# ══════════════════════════════════════════════════════════════════════════════

def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _setup_margins(doc)

    print("[1/9] Trang bìa và phụ bìa…")
    _add_cover(doc)
    _add_phubìa(doc)

    print("[2/9] Lời cảm ơn và lời cam đoan…")
    _add_loi_cam_on(doc)
    _add_loi_cam_doan(doc)

    print("[3/9] Tóm tắt khóa luận…")
    _add_abstract(doc)

    print("[4/9] Mục lục…")
    _add_muc_luc(doc)

    print("[5/9] Danh mục chữ viết tắt & bảng biểu…")
    _add_abbreviations(doc)
    _add_danh_muc_bang(doc)

    print("[6/9] Mở đầu…")
    _add_mo_dau(doc)

    print("[7/9] Chương 1–5 (nội dung chính)…")
    _add_chapters(doc)

    print("[8/9] Tài liệu tham khảo…")
    _add_references(doc)

    print("[9/9] Phụ lục…")
    _add_phu_luc(doc)

    # Thêm số trang footer
    _add_footer_page_number(doc)

    doc.save(str(OUT))
    print(f"\n✓ Hoàn tất: {OUT}")
    # Thống kê
    v = Document(str(OUT))
    print(f"  Đoạn văn: {len(v.paragraphs)} | Bảng: {len(v.tables)}")


if __name__ == "__main__":
    main()
