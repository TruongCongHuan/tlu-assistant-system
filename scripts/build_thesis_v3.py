#!/usr/bin/env python3
"""Tạo KLTN v3 — nội dung mới theo khung Khoa AI, giữ logo/hình từ bản v1.

Usage:
    python3 scripts/build_thesis_v3.py

Output:
    document/KLTN_TruongCongHuan_Final_v3.docx
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

REPO = Path(__file__).resolve().parent.parent
V1 = REPO / "document" / "KLTN_TruongCongHuan_Final (1).docx"
OUT = REPO / "document" / "KLTN_TruongCongHuan_Final_v3.docx"

sys.path.insert(0, str(REPO / "thesis" / "content_v3"))
import chapters as C  # noqa: E402

FONT = "Times New Roman"
BODY = Pt(13)
SECTION = Pt(14)
CHAPTER = Pt(16)

BODY_START = 175


def setup_page(doc: Document) -> None:
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(1.5)


def fmt_para(p, size=BODY, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, style=None):
    if style:
        try:
            p.style = style
        except KeyError:
            pass
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    if not p.runs and p.text:
        p.add_run(p.text)
    for run in p.runs:
        run.font.name = FONT
        run.font.size = size
        if bold:
            run.bold = True


def delete_range(doc: Document, start: int, end: int) -> None:
    to_remove = []
    for i in range(start, min(end, len(doc.paragraphs))):
        to_remove.append(doc.paragraphs[i]._element)
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def find_chapter(doc: Document, n: int) -> int | None:
    prefix = f"CHƯƠNG {n}"
    candidates = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.upper().startswith(prefix) and len(t) < 80:
            candidates.append(i)
    # Prefer body chapter (after danh mục), not TOC line
    for i in reversed(candidates):
        if i >= 150:
            return i
    return candidates[-1] if candidates else None


def find_refs(doc: Document) -> int | None:
    for i in range(150, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip().upper()
        if t == "TÀI LIỆU THAM KHẢO":
            return i
    return None


def find_heading(doc: Document, text: str, start: int = 0) -> int | None:
    for i in range(start, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip().upper().startswith(text.upper()):
            return i
    return None


def find_next_section(doc: Document, start: int) -> int:
    for i in range(start + 1, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip().upper()
        if t.startswith("DANH MỤC") or t.startswith("CHƯƠNG") or t.startswith("TÀI LIỆU"):
            return i
    return len(doc.paragraphs)


def replace_text_preserve_runs(p, new_text: str) -> None:
    if not p.runs:
        p.add_run(new_text)
        return
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        xml = r._element.xml
        if "blip" not in xml and "drawing" not in xml:
            r.text = ""


def _insert_after_el(doc: Document, ref_el, text: str, kind: str):
    new_p = OxmlElement("w:p")
    ref_el.addnext(new_p)
    para = Paragraph(new_p, doc)
    para.add_run(text)
    if kind == "chapter":
        fmt_para(para, CHAPTER, True, WD_ALIGN_PARAGRAPH.CENTER, "Heading 1")
    elif kind == "section":
        fmt_para(para, SECTION, True, WD_ALIGN_PARAGRAPH.LEFT, "Heading 2")
    elif kind == "subsection":
        fmt_para(para, BODY, True, WD_ALIGN_PARAGRAPH.LEFT, "Heading 3")
    elif kind == "bullet":
        fmt_para(para, BODY, style="List Bullet")
    else:
        fmt_para(para, BODY)
    return new_p


def _insert_table_after_el(doc: Document, ref_el, spec: dict):
    cap = spec.get("caption", "")
    if cap:
        ref_el = _insert_after_el(doc, ref_el, cap, "body")

    headers = spec["headers"]
    rows = spec["rows"]
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        tbl.style = "Table Grid"
    except KeyError:
        pass
    for ci, h in enumerate(headers):
        cell = tbl.rows[0].cells[ci]
        cell.text = h
        for cp in cell.paragraphs:
            for r in cp.runs:
                r.font.name = FONT
                r.font.size = Pt(11)
                r.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for cp in cell.paragraphs:
                for r in cp.runs:
                    r.font.name = FONT
                    r.font.size = Pt(11)

    tbl_el = tbl._element
    doc.element.body.remove(tbl_el)
    ref_el.addnext(tbl_el)
    return tbl_el


def update_cover_titles(doc: Document) -> None:
    titles = [
        "PHÁT TRIỂN HỆ THỐNG TRỢ LÝ AI HỌC VỤ",
        "TÍCH HỢP QUẢN LÝ SINH VIÊN",
    ]
    old = ["PHÁT TRIỂN ỨNG DỤNG QUẢN LÝ HỌC TẬP", "CÁ NHÂN TÍCH HỢP AI TƯ VẤN HỌC TẬP"]
    ti = 0
    for p in doc.paragraphs[:170]:
        t = p.text.strip().upper()
        for o in old:
            if o in t and ti < len(titles):
                replace_text_preserve_runs(p, titles[ti])
                ti += 1
                break


def update_abstract(doc: Document) -> None:
    idx = find_heading(doc, "TÓM TẮT KHÓA LUẬN")
    if idx is None:
        return
    end = find_heading(doc, "MỤC LỤC", idx + 1)
    if end is None:
        end = idx + 10
    # Xóa các đoạn tóm tắt cũ, chèn nội dung mới
    delete_range(doc, idx + 1, end)
    ref_el = doc.paragraphs[idx]._element
    for part in C.ABSTRACT.split("\n\n"):
        part = part.strip()
        if part:
            ref_el = _insert_after_el(doc, ref_el, part, "body")


def update_abbreviations(doc: Document) -> None:
    idx = find_heading(doc, "DANH MỤC CHỮ VIẾT TẮT", 120)
    if idx is None:
        return
    end = find_next_section(doc, idx)
    delete_range(doc, idx + 1, end)
    ref_el = doc.paragraphs[idx]._element
    for abbr, full in C.ABBREVIATIONS:
        ref_el = _insert_after_el(doc, ref_el, f"{abbr}: {full}", "body")


def update_table_list(doc: Document) -> None:
    idx = find_heading(doc, "DANH MỤC BẢNG BIỂU", 150)
    if idx is None:
        return
    end = find_next_section(doc, idx)
    delete_range(doc, idx + 1, end)
    ref_el = doc.paragraphs[idx]._element
    for line in ["Bảng 4.3. Phân loại bộ câu hỏi đánh giá RAG",
                 "Bảng 4.4. Kết quả ablation pipeline RAG"]:
        ref_el = _insert_after_el(doc, ref_el, line, "body")


def insert_v3_body(doc: Document) -> None:
    ch1 = find_chapter(doc, 1)
    refs = find_refs(doc)
    if ch1 is None or refs is None:
        print("  [ERROR] Không xác định được vùng nội dung cũ")
        sys.exit(1)

    insert_at = max(ch1 - 1, 0)
    anchor = doc.paragraphs[insert_at]._element

    delete_range(doc, ch1, refs)
    print(f"  ✓ Xóa nội dung cũ: đoạn {ch1}–{refs - 1}")

    table_map = {
        "EVAL_CATEGORY_TABLE": C.EVAL_CATEGORY_TABLE,
        "ABLACTION_TABLE": C.ABLACTION_TABLE,
    }

    ref_el = anchor
    for kind, text in C.BODY_BLOCKS:
        if kind == "table_ref":
            ref_el = _insert_table_after_el(doc, ref_el, table_map[text])
        else:
            ref_el = _insert_after_el(doc, ref_el, text, kind)

    ref_el = _insert_after_el(doc, ref_el, "TÀI LIỆU THAM KHẢO", "chapter")
    for ref in C.REFERENCES:
        ref_el = _insert_after_el(doc, ref_el, ref, "body")

    ref_el = _insert_after_el(doc, ref_el, "PHỤ LỤC", "chapter")
    ref_el = _insert_after_el(doc, ref_el, "Phụ lục A. Cấu trúc mã nguồn chính", "section")
    ref_el = _insert_after_el(doc, ref_el,
        "lib/services/local_rag_service.dart — RAG on-device. "
        "services/rag-api/app/rag.py — RAG server. "
        "services/rag-api/eval/run_eval.py — eval harness.", "body")
    ref_el = _insert_after_el(doc, ref_el, "Phụ lục B. Hướng dẫn triển khai", "section")
    _insert_after_el(doc, ref_el, "Xem README.md và docs/GIT_SETUP.md trong repository dự án.", "body")


def add_page_number_footer(doc: Document) -> None:
    try:
        for section in doc.sections:
            footer = section.footer
            footer.is_linked_to_previous = False
            para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            para.clear()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            run.font.name = FONT
            run.font.size = BODY
            fc_b = OxmlElement("w:fldChar")
            fc_b.set(qn("w:fldCharType"), "begin")
            run._r.append(fc_b)
            ir = para.add_run()
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = " PAGE "
            ir._r.append(instr)
            er = para.add_run()
            fc_e = OxmlElement("w:fldChar")
            fc_e.set(qn("w:fldCharType"), "end")
            er._r.append(fc_e)
    except Exception as exc:  # noqa: BLE001
        print(f"  [WARN] Số trang: {exc}")


def verify(doc: Document) -> None:
    counts = {}
    for i, p in enumerate(doc.paragraphs):
        if i < BODY_START:
            continue
        t = p.text.strip()
        for key in ["CHƯƠNG 1", "CHƯƠNG 5", "1.1.", "3.1.", "4.4.", "5.1."]:
            if t.startswith(key) and len(t) < 100:
                counts[key] = counts.get(key, 0) + 1
    for k, c in counts.items():
        flag = "OK" if c == 1 else f"CẢNH BÁO ×{c}"
        print(f"  {k}: {flag}")


def main() -> None:
    if not V1.exists():
        print(f"[ERROR] Không tìm thấy {V1}")
        sys.exit(1)

    print(f"[1/5] Copy {V1.name} → {OUT.name}")
    shutil.copy2(V1, OUT)
    doc = Document(str(OUT))

    print("[2/5] Cập nhật phần mở đầu…")
    update_cover_titles(doc)
    update_abstract(doc)
    update_abbreviations(doc)
    update_table_list(doc)

    print("[3/5] Ghi nội dung v3 (Chương 1–5 + TLTK + Phụ lục)…")
    insert_v3_body(doc)

    print("[4/5] Định dạng…")
    setup_page(doc)
    add_page_number_footer(doc)

    print("[5/5] Lưu & kiểm tra…")
    doc.save(str(OUT))
    v = Document(str(OUT))
    print(f"\nHoàn tất: {OUT}")
    print(f"  Đoạn văn: {len(v.paragraphs)} | Bảng: {len(v.tables)} | Hình: {len(v.inline_shapes)}")
    verify(v)
    print("\nTrong Word: Ctrl+A → F9 để cập nhật Mục lục / danh mục hình bảng.")


if __name__ == "__main__":
    main()
